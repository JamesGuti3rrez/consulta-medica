"""Informe Word de HU-014, condicionado a evidencias aprobadas y limpieza real confirmada."""
from collections import Counter
from pathlib import Path
import hashlib,json
import sys
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor
from dotenv import dotenv_values
from PIL import Image

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
sys.path.insert(0, str(ROOT / 'pruebas'))
from codigos_documentacion import align_document, destination as report_destination
OUT=report_destination('HU-014')


def xml_tests(path):
    root=ET.parse(path).getroot()
    assert not root.findall('.//failure') and not root.findall('.//error'),'XML con fallos'
    return list(root.iter('testcase'))


def main():
    specific=xml_tests(HERE/'pytest.xml');regression=xml_tests(HERE/'regresion.xml')
    sql=json.loads((HERE/'sql.json').read_text(encoding='utf-8'))
    browser=json.loads((HERE/'navegador.json').read_text(encoding='utf-8'))
    loading=json.loads((HERE/'carga_grupos.json').read_text(encoding='utf-8'))
    extra=json.loads((HERE/'ui_adicional.json').read_text(encoding='utf-8'))
    live=json.loads((HERE/'supabase_real.json').read_text(encoding='utf-8'))
    captures=json.loads((HERE/'capturas_terminal.json').read_text(encoding='utf-8'))
    assert all(data['resultado']=='PASA' for data in [sql,browser,loading,extra,live])
    assert live['limpieza']['confirmada']
    for row in captures:
        assert hashlib.sha256((ROOT/row['registro']).read_bytes()).hexdigest()==row['sha256']
    doc=Document();section=doc.sections[0]
    section.page_width,section.page_height=Cm(21),Cm(29.7)
    section.left_margin=section.right_margin=Cm(2);section.top_margin=section.bottom_margin=Cm(1.8)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.1
    for style,size in [('Title',27),('Heading 1',17),('Heading 2',12)]:
        doc.styles[style].font.name='Arial';doc.styles[style].font.size=Pt(size);doc.styles[style].font.color.rgb=RGBColor.from_string('176875')
    section.header.paragraphs[0].text='CONSULTA MÉDICA VIRTUAL  /  SPRINT 2  /  HU-014'
    section.header.paragraphs[0].style=doc.styles['Caption']
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('HU-014 · 30/10/2026 | Página ').font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    def p(t='',bold=False):
        para=doc.add_paragraph();para.add_run(t).bold=bold;return para
    def h(t,level=1):doc.add_heading(t,level)
    def page(t):doc.add_page_break();h(t)
    def bullet(t):doc.add_paragraph(t,style='List Bullet')
    def code(t):
        para=p();para.paragraph_format.line_spacing=1;run=para.add_run(t);run.font.name='Consolas';run.font.size=Pt(8)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for cell,text in zip(t.rows[0].cells,headers):
            cell.text=str(text);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4F3F2');cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:run.bold=True
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(4)
                    for run in para.runs:run.font.size=Pt(9)
        p()
    def picture(name,caption,max_height=17):
        file=HERE/'capturas'/name
        with Image.open(file) as image:w,height=image.size
        para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(file),width=Cm(min(16.4,max_height*w/height)))
        p(caption).style=doc.styles['Caption']

    doc.core_properties.title='HU-014 — Panel Doctor, grupos y desempeño'
    doc.core_properties.author='Equipo Consulta Médica Virtual'
    p('INFORME DE IMPLEMENTACIÓN Y VERIFICACIÓN',True);doc.add_heading('HU-014',0)
    h('Panel Doctor, grupos y desempeño');p('Sprint 2 · Refinamiento aprobado · 30 de octubre de 2026',True)
    p('Estado técnico: implementado y verificado, con migración aplicada y comprobación en Supabase real. La aceptación académica del Product Owner se mantiene como decisión independiente.')
    h('Historia de usuario',2)
    p('Como doctor responsable de un grupo de alumnos, quiero crear y administrar grupos, incorporar alumnos registrados, asignarles casos clínicos y consultar sus resultados y tiempos de resolución, para comparar su desempeño dentro del grupo y orientar su aprendizaje.')
    h('Resultado',2)
    bullet(f'{len(regression)} pruebas de regresión aprobadas, incluidas {len(specific)} específicas de API y métricas de grupos; tres pruebas vivo históricas deseleccionadas en la ejecución local.')
    bullet(f'{len(sql["casos"])} pruebas SQL sobre PostgreSQL local, ejecutando la migración real dos veces.')
    bullet(f'{len(browser["pruebas"])} comprobaciones de navegador, vista de 390 px y control de cambio de identidad.')
    bullet(f'{len(loading["pruebas"])} verificaciones adicionales de carga inicial, recuperación de grupos vacíos y rechazo de nombres duplicados.')
    bullet(f'{len(live["pruebas"])} comprobaciones con Supabase real; limpieza confirmada de siete cuentas, un grupo y cuatro sesiones con sus eventos.')
    p('El Panel Doctor conserva las funciones anteriores de catálogo, importación y evidencia. El nombre visible cambia; docente.html sigue funcionando. Mis grupos es la nueva entrada del alumno autenticado.')

    page('1. Alcance y recorrido funcional')
    table(['Actor','Acción','Resultado'],[
      ('Doctor','Crear o editar un grupo','Nombre y descripción persistentes, visibles a su propietario.'),
      ('Doctor','Añadir correo registrado','Integrante vinculado a Auth; repetición sin duplicado.'),
      ('Doctor','Asignar caso curado','Actividad con título sin solución y roster de alumnos.'),
      ('Alumno','Entrar en Mis grupos','Solo sus grupos y actividades; requiere sesión iniciada.'),
      ('Alumno','Iniciar, continuar o consultar entrega','Una sesión por alumno y actividad; recarga sin duplicados.'),
      ('Doctor','Comparar la misma actividad','Estado, diagnóstico entregado, acierto, tiempo y dos puntajes.'),
      ('Doctor','Retirar integrante / cerrar actividad','Revocación de acceso / bloqueo de nuevos inicios; historial conservado.'),
    ])
    p('Las actividades son independientes de la práctica libre y de una modalidad gamificada. No se añadieron premios ni una clasificación pública de compañeros. La autenticación obligatoria se aplica al trabajo de grupo.')
    h('Una entrega por actividad',2)
    p('La primera consulta se crea y vincula al grupo en una transacción. Un segundo inicio devuelve la misma sesión, aunque provenga de otra instancia. Si ya finalizó, la interfaz abre el resultado original. Para otro ejercicio o nueva entrega, el doctor crea una actividad nueva.')
    p('Los alumnos añadidos reciben las actividades abiertas en ese momento. Una actividad cerrada conserva el roster que ya tenía; reabrirla permite continuar a ese roster. Otra cohorte se trabaja mediante otra actividad.')

    page('2. Datos, transacciones y autorización')
    table(['Tabla','Función'],[
      ('doctor_grupos','Grupo y doctor propietario.'),('doctor_grupo_miembros','Cuenta del alumno y pertenencia activa/inactiva.'),
      ('doctor_asignaciones','Caso curado, título, grupo y apertura de la actividad.'),
      ('doctor_participaciones','Roster y vínculo único alumno–actividad–sesión.'),
      ('sma_sesiones / sma_eventos','Estado, calificación y eventos de la consulta, reutilizados sin duplicar notas.'),
    ])
    p('La migración 0005_grupos_doctor.sql requiere las tablas de 0002 y 0003. No depende de perfiles o intentos, ni de activar el historial opcional. El usuario confirmó la ejecución en SQL Editor y la prueba real verificó su disponibilidad.')
    p('La migración 0006_grupos_nombres_unicos.sql añade un índice único por doctor, ignorando mayúsculas y espacios repetidos. Normaliza el nombre al crear o editar. Conserva los grupos duplicados existentes y sus vínculos, añadiendo (2), (3), etc. sin colisionar con nombres existentes ni superar 100 caracteres. La base protege también solicitudes simultáneas. El usuario confirmó su aplicación y el rechazo 409 se comprobó en Supabase real.')
    p('El inicio bloquea el grupo, comprueba pertenencia y crea sesión, evento y vínculo en una transacción. La clave primaria alumno–actividad y el session_id único impiden duplicar la consulta. Un lote de asignaciones inválido revierte sus escrituras. El cierre usa una actualización condicional: solo la primera entrega modifica la sesión y emite su evento; un turno tardío no puede reabrirla.')
    h('Controles de acceso',2)
    for text in [
      'El token se verifica con Supabase Auth. El actor lo fija el servidor y no el cuerpo JSON del navegador.',
      'Profesor y administrativo pueden gestionar sus propios grupos. Un doctor ajeno recibe 404 sin participantes ni resultados.',
      'Solo integrantes activos consultan sus actividades. Un invitado recibe 401; un alumno no asignado recibe 404.',
      'La autorización de sesiones comprueba también la pertenencia de grupo en lectura, mensajes, eventos, reporte y cierre. Retirar al alumno revoca ese acceso.',
      'Las cuatro tablas tienen RLS activado y acceso revocado a anon/authenticated. Las RPC son exclusivas del backend con service_role y vuelven a validar rol/ownership.',
      'El título de una actividad se rechaza si contiene diagnóstico o sinónimos protegidos. Las vistas previas no muestran solución, diagnóstico entregado ni notas.',
      'La interfaz abandona el panel cuando cambia la identidad autenticada. Mis grupos limpia los datos al cerrar sesión o cambiar de usuario.',
    ]:bullet(text)

    page('3. Definición de métricas y comparación')
    table(['Indicador','Regla'],[
      ('Estado','Sin iniciar, en curso, finalizada o pendiente de revisión.'),
      ('Tiempo de resolución','finished_at menos started_at. Timestamps del servidor; incluye pausas.'),
      ('Acierto','Referencia o sinónimo correcto, incorrecto, vacío o ambiguo.'),
      ('Puntaje diagnóstico','Nota y motivo persistidos en diagnosis_evaluation, tras el cierre.'),
      ('Puntaje del proceso','evaluation.score de la consulta, independiente de la nota diagnóstica.'),
      ('Tasa de acierto','Correctos / entregas evaluadas. Ambiguas pendientes fuera del denominador.'),
      ('Promedios','Solo integrantes activos; tiempos de entregas finalizadas. Sin entregar conserva null.'),
    ])
    p('El tiempo no representa latencia de API/LLM ni tiempo activo frente a pantalla. Los alumnos que no han terminado no reciben un cero. Los integrantes retirados conservan su historial visible al doctor, pero quedan fuera de los promedios actuales.')
    p('Una respuesta ambigua mantiene su nota automática y motivo, con estado pendiente de revisión. No se implementó adjudicación manual ni una nota combinada: no se inventaron pesos entre diagnóstico y proceso.')
    h('Ejemplo obtenido en Supabase real',2)
    summary=live['comparativa_previa_al_retiro']['resumen']
    table(['Asignados','Iniciados','Finalizados','Pendientes de revisión','Evaluadas','Acierto'],[
      (summary['asignados'],summary['iniciados'],summary['finalizados'],summary['pendientes_revision'],summary['entregas_evaluadas'],str(summary['tasa_acierto'])+'%')])
    p('Cinco alumnos: correcto, incorrecto, ambiguo, en curso y sin iniciar. Tres entregas finalizadas y dos evaluadas; un acierto dividido entre dos = 50%. La ambigüedad no se convirtió en error definitivo y la falta de entrega no se convirtió en cero.')

    page('4. Evidencia visual: gestión del grupo')
    picture('miembros.png','Figura 1. El doctor incorpora alumnos registrados y puede retirarlos conservando su historial.',max_height=12)
    picture('asignar.png','Figura 2. Selección del caso y del título que verá el alumno. La evidencia de navegador comprueba la asignación efectiva; esta captura muestra el formulario.',max_height=9)
    page('5. Evidencia visual: alumno y resultado')
    picture('alumno-grupos.png','Figura 3. Mis grupos muestra únicamente actividades propias y acceso al caso asignado.',max_height=15)
    picture('alumno-resultado.png','Figura 4. El resultado del diagnóstico se muestra después del cierre; la calificación procede del backend.',max_height=5)
    page('6. Evidencia visual: seguimiento y comparación')
    picture('comparativa-en-curso.png','Figura 5. Se distingue al alumno en curso de los finalizados y no se publica su diagnóstico ni nota.',max_height=11)
    picture('comparativa-final.png','Figura 6. Comparación final con respuesta ambigua pendiente y tasa de acierto del 50% sobre dos entregas evaluadas.',max_height=11)
    p('Las capturas de navegador usan el frontend y API reales contra PostgreSQL local con la migración real. Auth, datos y dispositivos de voz son fixtures para reproducibilidad; no se presentan como cuentas productivas.')

    page('7. Pruebas automatizadas y comprobación remota')
    table(['Capa','Casos aprobados','Fuente'],[
      ('Regresión pytest',len(regression),'regresion.xml'),('API y métricas HU-014',len(specific),'pytest.xml'),
      ('SQL y privilegios',len(sql['casos']),'sql.json'),('Navegador',len(browser['pruebas']),'navegador.json'),
      ('Carga y nombres únicos',len(loading['pruebas']),'carga_grupos.json'),
      ('Supabase real',len(live['pruebas']),'supabase_real.json'),('Interfaz adicional','390 px y cambio de identidad','ui_adicional.json'),
    ])
    h('Escenarios relevantes',2)
    for text in [
      'Crear, editar, incorporar cuentas registradas y asignar casos; rechazo de alumno inexistente y de rol no permitido.',
      'Otro doctor, invitado y alumno ajeno denegados en lectura y escritura; intento de imponer dueño o caso desde el cliente.',
      'Inicio simultáneo y reintentos: una sesión y un evento inicial. Cierres concurrentes: misma nota y un solo evento de cierre. La recarga conserva consulta y resultado.',
      'Cierre incorrecto, correcto y ambiguo; notas independientes, valores null antes del cierre y denominadores correctos.',
      'Retirada de integrante revoca las cinco operaciones sobre su sesión y conserva la fila histórica.',
      'Actividad cerrada bloquea nuevos inicios y mantiene la consulta de quien ya comenzó.',
      'RLS y privilegios deniegan acceso directo de navegador a tablas/RPC, tanto localmente como en Supabase real.',
    ]:bullet(text)
    p('La prueba remota usa Supabase Auth, RPC, tablas, casos y sesiones reales. Un proveedor LLM local controlado evita gasto de generación y no sustituye ninguna regla de calificación. Se eliminaron y verificaron siete cuentas temporales, un grupo y cuatro sesiones con sus eventos. No se modificaron usuarios existentes, casos curados ni .env.')

    page('8. Capturas de los registros originales')
    p('Las siguientes imágenes son capturas nativas de ventanas propias de PowerShell que muestran los registros originales guardados al ejecutar las pruebas. No son salidas reconstruidas ni capturas de la terminal integrada de VS Code. Se conserva el SHA-256 de cada registro.')
    picture('terminal-regresion.png',f'Figura 7. Registro de la regresión: {len(regression)} pruebas aprobadas y tres deseleccionadas.',max_height=10)
    picture('terminal-sql.png',f'Figura 8. {len(sql["casos"])} pruebas de migraciones, transacciones, nombres únicos y permisos en PostgreSQL local.',max_height=10)
    page('8.1. Registro de Supabase real')
    picture('terminal-supabase_real.png','Figura 9. Verificación remota del flujo, ownership, métricas y limpieza de datos temporales.',max_height=18)

    page('8.2. Corrección: grupos vacíos y nombres duplicados')
    p('Se reprodujo una carrera: el perfil podía verificarse antes de cargarse el script de grupos; el evento de autorización se perdía. Ahora el módulo también consulta el perfil ya verificado. Si se abre antes de verificar la sesión, la carga se reanuda cuando llega el perfil. Los grupos se consultan independientemente del catálogo: un catálogo lento o temporalmente no disponible no bloquea la lista. Cuando llegan los casos se completa el selector sin reconstruir el formulario.')
    p('Las pruebas de navegador usan sesiones controladas y API/PostgreSQL locales reales. Cubren salida y regreso con la misma identidad, dos grupos vacíos sin crear otro, módulo tardío, apertura temprana, catálogo lento o fallido y conflictos al crear/editar. El rechazo de duplicados y la consulta del grupo vacío también se comprobaron en Supabase real con cuentas temporales.')
    picture('grupos-vacios.png','Figura 10. Los grupos vacíos aparecen automáticamente al volver al panel.',max_height=9)
    picture('grupo-duplicado.png','Figura 11. Un nombre equivalente al existente se rechaza sin crear otro grupo.',max_height=9)
    page('8.3. Registro de carga y duplicados')
    picture('terminal-carga_grupos.png',f'Figura 12. Registro original de las {len(loading["pruebas"])} comprobaciones adicionales de navegador.',max_height=18)

    page('9. Archivos, ejecución y decisiones')
    table(['Componente','Archivos'],[
      ('Esquema y RPC','backend/supabase/migrations/0005_grupos_doctor.sql; 0006_grupos_nombres_unicos.sql'),
      ('API y persistencia','backend/server/api/groups.py; persistence/groups.py'),
      ('Métricas','backend/server/groups.py'),('Integración de sesión','api/ai.py; orchestration/service.py; main.py'),
      ('Doctor','frontend/public/docente.html; js/grupos-doctor.js; css/grupos.css'),
      ('Alumno','frontend/public/grupos.html; js/grupos-alumno.js; js/perfil.js'),
      ('Simulación asignada','frontend/public/js/simulacion.js'),('Pruebas','backend/tests/groups; pruebas/HU-014'),
    ])
    h('Terminal de VS Code',2)
    code('npm install\nnpm run test:groups\nnpm run test:groups:sql\nnpm run test:sprint2 -- --junitxml=../pruebas/HU-014/regresion.xml\nnpm start')
    p('Doctor: http://localhost:3000/docente.html?seccion=grupos\nAlumno: http://localhost:3000/grupos.html')
    p('Las instrucciones para repetir navegador, servicio SQL de fixture, prueba remota y capturas están en pruebas/HU-014/README.md. PGlite es dependencia de desarrollo; no forma parte del frontend entregado. python-docx y Pillow se usan para documentación.')
    h('Límites explícitos',2)
    p('Una entrega por actividad; otras prácticas se crean como actividades nuevas. Actualización manual de resultados mediante Actualizar, sin prometer tiempo real. Duración transcurrida con pausas, sin cronómetro de actividad. No se añadieron invitaciones por correo, ranking público, adjudicación manual de ambiguos ni ponderación combinada de notas. La práctica libre conserva su acceso existente.')
    h('Integridad de la evidencia',2)
    code('SHA-256 migración: '+hashlib.sha256((ROOT/'backend/supabase/migrations/0005_grupos_doctor.sql').read_bytes()).hexdigest())
    code('SHA-256 migración 0006: '+hashlib.sha256((ROOT/'backend/supabase/migrations/0006_grupos_nombres_unicos.sql').read_bytes()).hexdigest())
    for c in captures:p(Path(c['registro']).name+': '+c['sha256'])
    p('Referencia del motor de pruebas: https://pglite.dev/docs/ (documentación oficial). El estado aprobado proviene de los XML/JSON incluidos; el generador rechaza fallos y limpieza no confirmada.')

    entry = align_document(doc, 'HU-014')
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    with ZipFile(OUT) as archive:
        assert archive.testzip() is None
        xml=''
        for n in archive.namelist():
            if n.endswith('.xml'):
                raw=archive.read(n);ET.fromstring(raw);xml+=raw.decode('utf-8')
        for key,value in dotenv_values(ROOT/'.env').items():
            if value and len(value)>=12 and any(part in key.upper() for part in ('KEY','SECRET','TOKEN','PASSWORD')):assert value not in xml,'Credencial detectada'
    reopened=Document(OUT);assert len(reopened.inline_shapes)==12
    result={'tarea':entry['codigo'],'codigo_evidencia':'HU-014','documento':str(OUT.relative_to(ROOT)),'regresion':len(regression),'pruebas_api':len(specific),
      'pruebas_sql':len(sql['casos']),'navegador':len(browser['pruebas']),'supabase_real':len(live['pruebas']),
      'carga_grupos':len(loading['pruebas']),
      'limpieza_confirmada':True,'capturas_word':len(reopened.inline_shapes),'zip_xml_validos':True,'credenciales_incluidas':False}
    (HERE/'resumen.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
