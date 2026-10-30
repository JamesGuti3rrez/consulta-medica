"""Capturas nativas de registros originales e informe del diálogo directo."""
from pathlib import Path
from zipfile import ZipFile
import ctypes
import hashlib
import importlib.util
import json
import subprocess
import time
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Cm,Pt
from dotenv import dotenv_values
from PIL import Image

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[1]
OUT=ROOT/'Entregables/Sprint2/TA-017_Dialogo_estudiante_medico_paciente.docx'


def main():
    counts={}
    for label in ('pytest','regresion'):
        tree=ET.parse(HERE/(label+'.xml')).getroot()
        assert not tree.findall('.//failure') and not tree.findall('.//error')
        counts[label]=len(list(tree.iter('testcase')))
    browser=json.loads((HERE/'navegador.json').read_text(encoding='utf8'))
    assert browser['resultado']=='PASA'
    spec=importlib.util.spec_from_file_location('terminal_helper',ROOT/'pruebas/Sprint2/fotografiar_terminal.py')
    helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)
    captures=[]
    for label in ('pytest','navegador'):
        source=HERE/(label+'.txt');title='DIALOGO-DIRECTO-'+label;hwnd=None
        proc=subprocess.Popen(['conhost.exe','powershell.exe','-NoLogo','-NoProfile','-ExecutionPolicy','Bypass',
          '-File',str(ROOT/'pruebas/Sprint2/mostrar_registro.ps1'),'-Registro',str(source),'-Titulo',title,'-Filas','32'],cwd=ROOT)
        try:
            for _ in range(100):
                hwnd=helper.find_window(title)
                if hwnd:break
                time.sleep(.1)
            if not hwnd:raise RuntimeError('No se encontró la ventana propia de PowerShell.')
            time.sleep(1.5);helper.user32.ShowWindow(hwnd,9)
            assert helper.user32.SetWindowPos(hwnd,helper.wintypes.HWND(-1),20,20,0,0,0x0041)
            helper.user32.SetForegroundWindow(hwnd);time.sleep(2)
            rect=helper.wintypes.RECT();helper.user32.GetWindowRect(hwnd,ctypes.byref(rect))
            ctypes.windll.dwmapi.DwmGetWindowAttribute(hwnd,9,ctypes.byref(rect),ctypes.sizeof(rect))
            target=HERE/'capturas'/('terminal-'+label+'.png')
            helper.ImageGrab.grab(bbox=(rect.left,rect.top,rect.right,rect.bottom),all_screens=True).save(target)
            captures.append({'registro':str(source.relative_to(ROOT)),'captura':str(target.relative_to(ROOT)),
                             'sha256':hashlib.sha256(source.read_bytes()).hexdigest()})
            print('[PASA] Captura nativa '+label,flush=True)
        finally:
            if hwnd:helper.user32.PostMessageW(hwnd,0x0010,0,0)
            try:proc.wait(timeout=4)
            except subprocess.TimeoutExpired:proc.terminate()
    (HERE/'capturas_terminal.json').write_text(json.dumps(captures,ensure_ascii=False,indent=2),encoding='utf8')
    doc=Document();section=doc.sections[0]
    section.page_width,section.page_height=Cm(21),Cm(29.7)
    section.left_margin=section.right_margin=Cm(2)
    section.top_margin=section.bottom_margin=Cm(1.8)
    doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
    section.header.paragraphs[0].text='CONSULTA MÉDICA VIRTUAL / SPRINT 2 / TA-017'
    section.footer.paragraphs[0].text='TA-017 · 30 de octubre de 2026'
    doc.core_properties.title='TA-017 — Separar diálogo estudiante–médico–paciente'
    doc.add_heading('TA-017',0);doc.add_heading('Separar diálogo estudiante–médico–paciente',1)
    doc.add_paragraph('Implementación y verificación · 30 de octubre de 2026')
    doc.add_paragraph('Fuente: Backlog_Sprint2_Fases_y_Dependencias_Completado.docx. Responsable: Rodrigo Nicolas Flores Alvarez.')
    doc.add_paragraph('Resultado: el doctor presenta el caso y después entrega hallazgos o aclara solicitudes. El estudiante conversa directamente con el paciente o el servicio solicitado, sin un turno intermedio del doctor que repita su pregunta.')
    doc.add_heading('Criterios y desarrollo',1)
    for text in [
      'Diez intervenciones con emisor y destinatario distinguibles; sin repetir mecánicamente preguntas ni mostrar diagnóstico protegido.',
      'El reformulador se retiró del grafo activo y de la construcción del servicio. Recuperación, autorización y agente usan la pregunta original.',
      'Cada turno nuevo conserva estudiante y agente en el historial; doctor_question se devuelve como null para mantener compatibilidad del contrato.',
      'La interfaz omite preguntas intermediadas antiguas al recuperar actividades y no las reproduce por voz. El contexto del agente tampoco recibe esos ecos; los registros históricos no se borran.',
      'La espera de respuesta es neutral. El médico sigue entregando hallazgos de exploración física y aclarando solicitudes genéricas; laboratorio, radiología y especialista conservan su emisor.',
    ]:doc.add_paragraph(text,style='List Bullet')
    doc.add_heading('Resultados verificados',1)
    doc.add_paragraph(f'{counts["regresion"]} pruebas de regresión aprobadas y {counts["pytest"]} comprobaciones específicas aprobadas. Estas últimas cubren diez turnos, uso de la pregunta original, compatibilidad del historial y bloqueo de una respuesta protegida del paciente.')
    doc.add_paragraph(f'{len(browser["pruebas"])} comprobaciones de navegador: presentación inicial, diez turnos de dos mensajes y una lectura por voz, recuperación de una actividad antigua sin eco y ausencia de errores JavaScript.')
    doc.add_paragraph('Entorno: FastAPI y PostgreSQL local PGlite reales; Auth y dispositivos de voz controlados. El evento histórico es un dato sintético insertado exclusivamente en el fixture. No se ejecutaron escrituras de prueba en Supabase ni llamadas a Groq.')

    def picture(name,caption,height=15):
        with Image.open(HERE/'capturas'/name) as im:w,h=im.size
        doc.add_picture(str(HERE/'capturas'/name),width=Cm(min(16.4,height*w/h)))
        doc.add_paragraph(caption,style='Caption')
    doc.add_page_break();doc.add_heading('Evidencia de la conversación',1)
    picture('pregunta-directa.png','Figura 1. Después de la presentación, la pregunta del estudiante recibe directamente la respuesta del paciente.',height=12)
    picture('hallazgos-aclaracion.png','Figura 2. El doctor interviene para entregar hallazgos y aclarar solicitudes; no repite preguntas.',height=12)
    doc.add_page_break();doc.add_heading('Compatibilidad con actividades anteriores',1)
    picture('conversacion-recuperada.png','Figura 3. Al recuperar una actividad se conservan preguntas y respuestas, omitiendo el eco antiguo del doctor.')
    doc.add_paragraph('La recuperación probada es la existente de actividades de grupo. Este cambio no implementa la historia general de retomar sesiones interrumpidas.')
    doc.add_page_break();doc.add_heading('Registros de ejecución',1)
    doc.add_paragraph('Capturas nativas de ventanas propias de PowerShell que muestran los registros originales; no son capturas de la terminal integrada de VS Code. Sus huellas SHA-256 se conservan en capturas_terminal.json.')
    picture('terminal-pytest.png','Figura 4. Cuatro pruebas específicas aprobadas.',height=10)
    picture('terminal-navegador.png','Figura 5. Comprobaciones de mensajes, voz y recuperación.',height=10)
    doc.add_heading('Archivos y reproducción',1)
    doc.add_paragraph('backend/server/api/ai.py; orchestration/star_graph.py y service.py; ai/agents/prompts/paciente.md y medico_asistente.md; frontend/public/js/simulacion.js; backend/tests/integration/test_direct_dialogue.py.')
    doc.add_paragraph('Instrucciones completas: pruebas/TA-017_Dialogo/README.md. Ejecutar npm run test:sprint2 para la regresión con los nuevos escenarios incluidos.')
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    with ZipFile(OUT) as z:
        assert z.testzip() is None
        xml=''
        for name in z.namelist():
            if name.endswith('.xml'):
                raw=z.read(name);ET.fromstring(raw);xml+=raw.decode('utf8')
        for key,value in dotenv_values(ROOT/'.env').items():
            if value and len(value)>=12 and any(s in key.upper() for s in ('KEY','SECRET','TOKEN','PASSWORD')):assert value not in xml
    assert len(Document(OUT).inline_shapes)==5
    result={'tarea':'TA-017','documento':str(OUT.relative_to(ROOT)),'regresion':counts['regresion'],
            'pruebas_especificas':counts['pytest'],'navegador':len(browser['pruebas']),'capturas':5,'resultado':'PASA'}
    (HERE/'resumen.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
