"""Alineación documental con el backlog aprobado, conservando la evidencia original."""
from __future__ import annotations

import json
from pathlib import Path
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
METADATA = ROOT / 'docs/development/correspondencia_backlog_sprint2.json'
DATA = json.loads(METADATA.read_text(encoding='utf-8'))

# Referencias del backlog anterior. Se sustituyen de una vez, nunca en cascada.
REFERENCES = {
    'TA-017': 'TA-011', 'TA-007': 'TA-012', 'TA-027': 'TA-014',
    'HU-004': 'HU-005', 'EN-013': 'EN-005', 'TA-026': 'TA-015',
    'TA-028': 'TA-016', 'TA-029': 'TA-022', 'TA-019': 'TA-023',
    'HU-005': 'HU-006', 'HU-006': 'HU-007', 'HU-013': 'HU-010',
    'HU-012': 'HU-008', 'HU-011': 'HU-009', 'EN-012': 'EN-006',
    'HU-014': 'HU-011',
}
IDENTIFIER = re.compile(r'\b(?:TA|HU|EN)-\d{3}\b')
EVIDENCE_PATH = re.compile(r'''(?:pruebas|backend|frontend|data|docs|Entregables)[/\\][^\s"'`;,)]*''')


def task(legacy_id):
    return DATA['tareas'][legacy_id]


def destination(legacy_id):
    return ROOT / 'Entregables/Sprint2' / task(legacy_id)['archivo']


def remap(text, legacy_id=None):
    mapping = dict(REFERENCES)
    if legacy_id == 'TA-012':
        # El informe histórico de calificación usó TA-012, distinto del guardrail actual.
        mapping['TA-012'] = 'TA-021'
    protected = [m.span() for m in EVIDENCE_PATH.finditer(text)]

    def replace(match):
        if any(start <= match.start() < end for start,end in protected):
            return match.group()  # Rutas reales de pruebas/capturas conservadas.
        return mapping.get(match.group(), match.group())
    return IDENTIFIER.sub(replace, text)


def _set_text(paragraph, text):
    nodes = paragraph.xpath('.//w:t')
    if not nodes:
        run = OxmlElement('w:r'); node = OxmlElement('w:t')
        run.append(node); paragraph.append(run); nodes = [node]
    nodes[0].text = text
    nodes[0].set(qn('xml:space'), 'preserve')
    for node in nodes[1:]:
        node.text = ''


def _remap_paragraph(paragraph, legacy_id):
    nodes = paragraph.xpath('.//w:t')
    texts = [node.text or '' for node in nodes]
    text = ''.join(texts)
    offsets = []; position = 0
    for value in texts:
        offsets.append((position, position+len(value))); position += len(value)
    protected = [m.span() for m in EVIDENCE_PATH.finditer(text)]
    for match in reversed(list(IDENTIFIER.finditer(text))):
        if any(start <= match.start() < end for start,end in protected):
            continue
        replacement = remap(match.group(), legacy_id)
        if replacement == match.group():
            continue
        start = next(i for i, (a,b) in enumerate(offsets) if a <= match.start() < b)
        end = next(i for i, (a,b) in enumerate(offsets) if a <= match.end()-1 < b)
        prefix = (nodes[start].text or '')[:match.start()-offsets[start][0]]
        suffix = (nodes[end].text or '')[match.end()-offsets[end][0]:]
        nodes[start].text = prefix + replacement + (suffix if start == end else '')
        for node in nodes[start+1:end]:
            node.text = ''
        if start != end:
            nodes[end].text = suffix


def align_document(doc, legacy_id):
    """Se ejecuta una vez sobre el documento del generador, antes de guardarlo."""
    entry = task(legacy_id)
    for part in doc.part.package.parts:
        element = getattr(part, '_element', None)
        if element is not None:
            for paragraph in element.xpath('.//w:p'):
                _remap_paragraph(paragraph, legacy_id)

    for paragraph in doc.paragraphs:
        text = paragraph.text
        if paragraph.style.name == 'Title':
            _set_text(paragraph._p, entry['codigo'])
        elif text.startswith('Responsable según backlog:'):
            _set_text(paragraph._p, 'Responsable según backlog: ' + entry['responsable'])
        elif text.startswith('Sprint 2 · Product Backlog'):
            _set_text(paragraph._p, 'Sprint 2 · Backlog por fases y dependencias · 30 de octubre de 2026')
        elif text.startswith('Alcance acordado:'):
            _set_text(paragraph._p, 'Alcance documentado: los cinco puntos implementados conservan su alcance funcional. Las exclusiones del encargo inicial corresponden ahora a HU-006 (proponer y registrar diagnóstico) y TA-023 (métricas RAG y RAGAS). La renumeración no modifica resultados ni implica completar otras tareas.')
        elif text.startswith('La protección de sesiones ya estaba implementada'):
            _set_text(paragraph._p, 'La protección de sesiones se identifica como TA-011 en el backlog actual. Las pruebas adicionales se conservan en la carpeta histórica TA-017, sin cambiar sus resultados ni sus capturas.')
        elif text.startswith('Documento fuente:'):
            _set_text(paragraph._p, 'Documento fuente de códigos: ' + DATA['fuente'])
        elif text.startswith('SHA-256 del backlog:'):
            _set_text(paragraph._p, 'SHA-256 del backlog: ' + DATA['sha256_fuente'])
        elif text == 'Calificación del diagnóstico final':
            _set_text(paragraph._p, entry['titulo'])

    if not entry.get('incorporacion_manual_pendiente'):
        # Se toma del nuevo documento, sin volver a renumerar sus dependencias actuales.
        for index, paragraph in enumerate(doc.paragraphs):
            if paragraph.text == 'Criterios de aceptación del documento fuente':
                _set_text(doc.paragraphs[index+1]._p, entry['criterios'])

    note = (
        'Códigos actualizados según ' + DATA['fuente'] + ' (30/10/2026). '
        'Las capturas, registros, comandos y rutas de evidencia conservan los identificadores usados al ejecutar las pruebas; no se han alterado sus resultados. '
        f'Correspondencia de este informe: {legacy_id} → {entry["codigo"]}.'
    )
    if entry.get('incorporacion_manual_pendiente'):
        note = f'{entry["codigo"]} identifica el Panel Doctor por indicación del usuario, quien lo incorporará manualmente al backlog. Su código anterior era {legacy_id}; las capturas y rutas históricas conservan ese identificador. La renumeración no implica que HU-010 (rondas gamificadas) esté implementada.'
    paragraphs = doc.paragraphs
    anchor = next(p for p in paragraphs if p.style.name.startswith('Heading'))
    added = anchor.insert_paragraph_before(note)
    added.style = doc.styles['Caption']
    if 'responsable' in entry:
        metadata = f'Responsable según backlog: {entry["responsable"]}. {entry["numero_fase"]}: {entry["fase"]}. Prioridad {entry["prioridad"]}; estimación {entry["sp"]}. Dependencias: {entry["dependencias"]}.'
        anchor.insert_paragraph_before(metadata)

    doc.core_properties.title = entry['codigo'] + ' — ' + entry['titulo']
    doc.core_properties.subject = 'Implementación y evidencia; códigos del backlog por fases y dependencias'
    doc.core_properties.keywords = entry['codigo'] + ', Sprint 2, documentación, evidencia'
    return entry
