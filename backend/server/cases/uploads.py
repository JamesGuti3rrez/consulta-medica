"""Lectura de JSON canónico y Excel tabular, sin inventar información clínica."""
from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from jsonschema import Draft202012Validator

from server.cases.rules import check_batch_rules, check_case_rules, check_case_warnings
from server.cases.validator import load_schema

MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_CASES = 100
MAX_ROWS = 5000
COLUMNS = {
    "Casos": ["case_id", "titulo", "especialidad", "nombre", "edad", "sexo", "motivo_consulta",
              "diagnostico", "sinonimos", "tratamiento", "acciones"],
    "Evidencias": ["case_id", "clave", "categoria", "etiqueta", "texto", "propietario", "tema"],
    "Diferenciales": ["case_id", "nombre", "imprescindible", "evidencias_descartan"],
    "Evaluacion": ["case_id", "tema", "evidencias", "peso", "critico", "error_si_falta"],
    "Examenes": ["case_id", "recomendados", "opcionales", "innecesarios"],
}
MODES = {"paciente": "al_preguntar", "medico_asistente": "al_examinar",
         "laboratorio": "al_solicitar_examen", "radiologia": "al_solicitar_examen"}


class UploadError(ValueError):
    pass


def split_values(value) -> list[str]:
    return [part.strip() for part in str(value or "").split("|") if part.strip()]


def _bool(value) -> bool:
    if value in (True, 1, "1", "true", "si", "sí", "Sí", "SI"):
        return True
    if value in (False, 0, None, "", "0", "false", "no", "No", "NO"):
        return False
    raise UploadError("Los valores de sí/no deben ser sí, no, true o false.")


def _integer(value, label: str) -> int:
    try:
        number = int(value)
        if isinstance(value, float) and number != value:
            raise ValueError()
        return number
    except (ValueError, TypeError, OverflowError) as exc:
        raise UploadError(f"{label}: se esperaba un número entero.") from exc


def read_upload(filename: str, content: bytes) -> list[dict]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise UploadError("El archivo está vacío o supera el límite de 4 MB.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (ValueError, UnicodeError, RecursionError) as exc:
            raise UploadError("El JSON no es válido. Guárdalo en UTF-8 y revisa su sintaxis.") from exc
        if isinstance(payload, dict) and "casos" in payload:
            payload = payload["casos"]
        cases = payload if isinstance(payload, list) else [payload]
    elif suffix == ".xlsx":
        cases = _read_excel(content)
    else:
        raise UploadError("Formatos admitidos: JSON y Excel .xlsx. Descarga una plantilla del panel.")
    if not cases or len(cases) > MAX_CASES or any(not isinstance(case, dict) for case in cases):
        raise UploadError("El archivo debe contener entre 1 y 100 objetos de caso.")
    return cases


def _read_excel(content: bytes) -> list[dict]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > 500 or sum(item.file_size for item in files) > 25 * 1024 * 1024:
                raise UploadError("El Excel descomprimido es demasiado grande.")
            if any("vbaproject" in item.filename.lower() for item in files):
                raise UploadError("Usa un archivo .xlsx sin macros.")
    except BadZipFile as exc:
        raise UploadError("El archivo no es un Excel .xlsx válido.") from exc
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UploadError("Falta el lector de Excel. Ejecuta npm run setup -- --extras dev.") from exc
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        tables = {}
        for name, columns in COLUMNS.items():
            if name not in workbook.sheetnames:
                raise UploadError(f"Falta la hoja {name}. Usa la plantilla del panel.")
            sheet = workbook[name]
            # No confiar en dimensiones proporcionadas por el archivo.
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            header = next(rows, ())
            if [cell.value for cell in header] != columns:
                raise UploadError(f"Las columnas de {name} deben ser: {', '.join(columns)}.")
            entries = []
            for number, cells in enumerate(rows, start=2):
                if number > MAX_ROWS:
                    raise UploadError(f"La hoja {name} supera las {MAX_ROWS} filas.")
                if len(cells) > len(columns) and any(c.value is not None for c in cells[len(columns):]):
                    raise UploadError(f"Hoja {name}, fila {number}: hay columnas adicionales.")
                if any(cell.data_type == "f" for cell in cells):
                    raise UploadError(f"Hoja {name}, fila {number}: reemplaza las fórmulas por valores.")
                values = [cell.value for cell in cells[:len(columns)]]
                if not any(value is not None and value != "" for value in values):
                    continue
                values += [None] * (len(columns) - len(values))
                entries.append(dict(zip(columns, values)))
            tables[name] = entries
        workbook.close()
    except UploadError:
        raise
    except Exception as exc:
        raise UploadError("No se pudo leer el Excel. Usa la plantilla y revisa que no esté cifrado.") from exc
    cases, known_ids = [], set()
    declared = [str(row["case_id"] or "").strip() for row in tables["Casos"]]
    if len(set(declared)) != len(declared) or any(not case_id for case_id in declared):
        raise UploadError("Cada fila de Casos debe tener un case_id único.")
    for name in COLUMNS:
        if name != "Casos" and any(row["case_id"] not in declared for row in tables[name]):
            raise UploadError(f"La hoja {name} contiene filas de casos que no aparecen en Casos.")
    for row in tables["Casos"]:
        case_id = str(row["case_id"] or "").strip()
        if not case_id or case_id in known_ids:
            raise UploadError("Cada fila de Casos debe tener un case_id único.")
        known_ids.add(case_id)
        evidence_rows = [item for item in tables["Evidencias"] if item["case_id"] == case_id]
        evidence, ids = [], {}
        for item in evidence_rows:
            key = str(item["clave"] or "").strip()
            if not key or key in ids:
                raise UploadError(f"{case_id}: cada evidencia necesita una clave única.")
            owner = str(item["propietario"] or "")
            if owner not in MODES:
                raise UploadError(f"{case_id}: propietario de evidencia no válido.")
            kind = {"laboratorio": "lab", "imagen": "imagen", "examen_fisico": "examen",
                    "signo_vital": "vital"}.get(item["categoria"], "anamnesis")
            evidence_id = f"{case_id}.{kind}.{key}"
            ids[key] = evidence_id
            evidence.append({"evidence_id": evidence_id, "category": item["categoria"],
                             "label": str(item["etiqueta"] or ""), "text": str(item["texto"] or ""),
                             "reveal_policy": {"owner": owner, "mode": MODES[owner],
                                               "requires": split_values(item["tema"])}})
        def refs(value):
            keys = split_values(value)
            if any(key not in ids for key in keys):
                raise UploadError(f"{case_id}: una evaluación o diferencial cita una clave de evidencia inexistente.")
            return [ids[key] for key in keys]
        differentials = [
            {"name": str(item["nombre"] or ""), "must_not_miss": _bool(item["imprescindible"]),
             "discard_evidence_ids": refs(item["evidencias_descartan"])}
            for item in tables["Diferenciales"] if item["case_id"] == case_id
        ]
        questions = []
        for item in tables["Evaluacion"]:
            if item["case_id"] == case_id:
                questions.append({"id": f"kq{len(questions)+1}", "topic": str(item["tema"] or ""),
                                  "evidence_ids": refs(item["evidencias"]),
                                  "weight": _integer(item["peso"], "Peso"), "critical": _bool(item["critico"]),
                                  "error_if_missing": str(item["error_si_falta"] or "anamnesis_incompleta")})
        tests = [item for item in tables["Examenes"] if item["case_id"] == case_id]
        if len(tests) != 1:
            raise UploadError(f"{case_id}: necesita exactamente una fila en Examenes.")
        cases.append({
            "schema_version": "1.1.0", "case_id": case_id, "title": str(row["titulo"] or ""),
            "specialty": row["especialidad"], "language": "es", "status": "borrador",
            "source": {"dataset": "manual", "item_id": f"docente/{case_id}", "license": "Aportado por el docente para uso educativo", "method": "manual"},
            "public": {
                "patient": {"alias": str(row["nombre"] or ""), "age_years": _integer(row["edad"], "Edad"), "sex": row["sexo"]},
                "chief_complaint": {"evidence_id": f"{case_id}.motivo.queja_principal", "category": "motivo",
                                   "label": "Motivo de consulta", "text": str(row["motivo_consulta"] or "")},
                "items": [{"evidence_id": f"{case_id}.anamnesis.contexto_atencion", "category": "anamnesis",
                           "label": "Contexto de atención", "text": "Acude a valoración por el motivo de consulta."}],
                "allergies_recorded": False,
            },
            "on_request": evidence,
            "protected": {
                "diagnosis": {"name": str(row["diagnostico"] or ""), "synonyms": split_values(row["sinonimos"])},
                "differentials": differentials,
                "treatment_reference": {"summary": str(row["tratamiento"] or ""), "key_actions": split_values(row["acciones"])},
                "contraindications": [],
                "evaluation": {"key_questions": questions, "tests": {
                    "recommended": split_values(tests[0]["recomendados"]), "optional": split_values(tests[0]["opcionales"]),
                    "unnecessary": split_values(tests[0]["innecesarios"])}}},
            "provenance": {"created_by": "Carga Excel docente", "created_at": datetime.now(timezone.utc).date().isoformat(),
                           "validator_version": "1.1.0", "validated_at": None, "curated_by": None,
                           "curated_at": None, "rejection_reason": None},
        })
    for name in COLUMNS:
        if name != "Casos" and any(row["case_id"] not in known_ids for row in tables[name]):
            raise UploadError(f"La hoja {name} contiene filas de casos que no aparecen en Casos.")
    return cases


def clinical_content(case: dict) -> dict:
    return {key: value for key, value in case.items() if key not in {"provenance", "status"}}


def prepare_cases(cases: list[dict], existing: list[dict], user_id: str) -> tuple[list[dict], list[dict], list[str]]:
    if not cases or len(cases) > MAX_CASES or any(not isinstance(case, dict) for case in cases):
        raise UploadError("El lote debe tener entre 1 y 100 casos.")
    prepared, errors, warnings = [], [], []
    remote = {case["case_id"]: case for case in existing}
    today = datetime.now(timezone.utc).date().isoformat()
    validator = Draft202012Validator(load_schema())
    for number, raw in enumerate(cases, start=1):
        case = json.loads(json.dumps(raw, ensure_ascii=False))
        case_id = case.get("case_id", f"fila {number}")
        previous = remote.get(case_id) if isinstance(case_id, str) else None
        if case.get("status") == "rechazado":
            errors.append({"case_id": str(case_id), "path": "status", "message": "Un caso rechazado requiere revisión antes de publicarse."})
            continue
        provenance = case.get("provenance") if isinstance(case.get("provenance"), dict) else {}
        case["status"] = "curado"
        case["provenance"] = {
            "created_by": provenance.get("created_by") or f"docente:{user_id}",
            "created_at": provenance.get("created_at") or today, "validator_version": "1.1.0",
            "validated_at": today, "curated_by": f"docente:{user_id}", "curated_at": today, "rejection_reason": None,
        }
        schema_errors = list(validator.iter_errors(case))
        for error in schema_errors[:10]:
            errors.append({"case_id": str(case_id), "path": ".".join(map(str, error.path)),
                           "message": error.message[:350]})
        if schema_errors:
            continue
        for error in check_case_rules(case):
            errors.append({"case_id": case_id, "path": "", "message": error})
        warnings += [f"{case_id}: {warning}" for warning in check_case_warnings(case)]
        if previous:
            if clinical_content(previous) != clinical_content(case):
                errors.append({"case_id": case_id, "path": "case_id",
                               "message": "Ya existe otro contenido con este ID. Asigna un ID nuevo; no se sobrescribirá."})
            else:
                case = previous
        prepared.append(case)
    if not errors:
        # Detectar IDs y fuentes repetidos tanto en el lote como frente al catálogo existente.
        errors += [{"case_id": "", "path": "", "message": message} for message in check_batch_rules(prepared)
                   if not message.startswith("R18")]
        combined = {**remote, **{case["case_id"]: case for case in prepared}}
        errors += [{"case_id": "", "path": "", "message": message}
                   for message in check_batch_rules(list(combined.values()))]
    return prepared, errors, warnings

